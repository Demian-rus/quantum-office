from pathlib import Path
from shutil import which
from io import BytesIO
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
root=Path(__file__).resolve().parents[1]
with sync_playwright() as p:
 browser=p.chromium.launch(executable_path=which('chromium') or None,headless=True,args=['--no-sandbox','--disable-dev-shm-usage'])
 page=browser.new_page()
 page.set_content((root/'dist/index.html').read_text(),wait_until='load')
 data=page.evaluate('''async()=>{const b=newWorkbook(),s=currentSheet(b);setCell(s,'B2','Недельный график');mergeCells(s,'B2','D4');applyBorderPreset(s,'B2','D4','outline');setStyle(s,'B2',{bg:'#eaf5de',wrap:true});return Array.from(await exportXLSX(b));}''')
 wb=load_workbook(BytesIO(bytes(data)));sh=wb.active
 assert 'B2:D4' in [str(m) for m in sh.merged_cells.ranges]
 assert sh['B2'].value=='Недельный график'
 assert sh['B2'].fill.fgColor.rgb=='FFEAF5DE'
 assert sh['B2'].alignment.wrap_text is True
 assert sh['B2'].border.top.style=='thin'
 # A merge-safe XLSX round trip: read the OOXML back into Quantum.
 imported=page.evaluate('''async bytes=>{const b=await importXLSX(Uint8Array.from(bytes),'test.xlsx');return {merges:b.sheets[0].merges,style:b.sheets[0].cells.B2.style};}''',data)
 assert imported['merges']==[{'from':'B2','to':'D4'}],imported
 assert imported['style']['borders']['top']==True,imported
 print('PASS merge, border, color, wrap XLSX export-openpyxl-import')
 browser.close()
