from pathlib import Path
import io
from shutil import which
from playwright.sync_api import sync_playwright
from openpyxl import load_workbook
root=Path(__file__).resolve().parents[1]
with sync_playwright() as p:
    browser=p.chromium.launch(**({'executable_path':which('chromium')} if which('chromium') else {}),headless=True,args=['--no-sandbox','--disable-dev-shm-usage'])
    page=browser.new_page(viewport={'width':1366,'height':768})
    page.set_content((root/'dist/index.html').read_text(),wait_until='load')
    result=page.evaluate('''async()=>{
      const b=newWorkbook('Список'),s=currentSheet(b);
      [['Категория','Сумма'],['яблоко','100'],['груша','200'],['слива','300'],['яблоко','400']].forEach((row,r)=>row.forEach((v,c)=>setCell(s,cellName(r,c),v)));
      setAutoFilter(s,'A1','B5');setFilterValue(s,0,['яблоко','слива']);
      const bytes=await exportXLSX(b);
      const copy=await importXLSX(bytes,'Список.xlsx');
      return {bytes:Array.from(bytes),rules:copy.sheets[0].autoFilter.rules,visible:[1,2,3,4].filter(r=>filterRowVisible(copy.sheets[0],r))};
    }''')
    assert result['rules']=={'0':['яблоко','слива']}
    assert result['visible']==[1,3,4]
    wb=load_workbook(io.BytesIO(bytes(result['bytes'])))
    sh=wb.active
    assert sh.auto_filter.ref=='A1:B5'
    cols=sh.auto_filter.filterColumn
    assert len(cols)==1
    assert [v.val if hasattr(v,'val') else v for v in cols[0].filters.filter]==['яблоко','слива']
    assert sh.row_dimensions[3].hidden is True
    assert sh['A3'].value=='груша'
    print('PASS XLSX multi-value filter roundtrip')
    browser.close()
