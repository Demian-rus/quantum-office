"""TASK040 — Excel sheet dimensions, narrow columns, font point sizes and borders."""
from io import BytesIO
from pathlib import Path
from shutil import which
from tempfile import TemporaryDirectory
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Border,Side,Font,PatternFill,Alignment
from playwright.sync_api import sync_playwright

root=Path(__file__).resolve().parents[1]
with TemporaryDirectory() as tmp:
 path=Path(tmp)/'format-and-widths.xlsx'
 book=Workbook();ws=book.active;ws.title='График'
 ws.sheet_format.baseColWidth=8
 ws.sheet_format.defaultRowHeight=15
 ws.column_dimensions['A'].width=6.0
 ws.column_dimensions['B'].width=3.0
 ws.column_dimensions['C'].width=22.0
 ws.column_dimensions['D'].width=10.0
 ws.row_dimensions[4].height=12.75
 edge=Side(style='thin',color='FF000000')
 ws['A1']='Даты';ws['A1'].font=Font(name='Calibri',sz=14,bold=True)
 ws['A1'].fill=PatternFill(fill_type='solid',fgColor='FF00FF00')
 ws['A1'].border=Border(left=edge,right=edge,top=edge,bottom=edge)
 ws.merge_cells('A1:D1')
 ws['B3']='Узкий';ws['B3'].font=Font(name='Calibri',sz=10)
 ws['C4']='Расписание';ws['C4'].font=Font(name='Calibri',sz=13,italic=True)
 ws['C4'].border=Border(bottom=edge)
 ws['C4'].alignment=Alignment(horizontal='center')
 book.save(path)
 with sync_playwright() as p:
  browser=p.chromium.launch(executable_path=which('chromium') or None,headless=True,args=['--no-sandbox','--disable-dev-shm-usage'])
  page=browser.new_page(viewport={'width':1300,'height':900},accept_downloads=True)
  page.set_default_timeout(25000)
  errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
  page.set_content((root/'dist/index.html').read_text(encoding='utf-8'),wait_until='load')
  page.locator('#file-picker').set_input_files(str(path))
  page.wait_for_selector('.grid-cell[data-ref="A1"]')
  data=page.evaluate("""()=>{
   const s=sheet(),geo=gridGeometry(s),a=getEffectiveStyle(s,'A1'),c=getEffectiveStyle(s,'C4');
   const aCell=document.querySelector('.grid-cell[data-ref="A1"]');
   const cCell=document.querySelector('.grid-cell[data-ref="C4"]');
   return {defaultCol:s.defaultColumnWidth,defaultRow:s.defaultRowHeight,
    widths:Array.from({length:5},(_,i)=>geo.x[i+1]-geo.x[i]),
    height:rowHeight(s,3),merged:s.merges.some(m=>m.from==='A1'&&m.to==='D1'),
    a:a,c:c,
    aCss:{font:getComputedStyle(aCell).fontSize,border:getComputedStyle(aCell).borderTopColor},
    cCss:{font:getComputedStyle(cCell).fontSize,border:getComputedStyle(cCell).borderBottomColor},
    persist:decodeBook(encodeBook(book())).sheets[0].defaultColumnWidth};
  }""")
  assert data['defaultCol']==61 and data['defaultRow']==20,data
  assert data['widths'][:4]==[47,26,159,75],data
  assert data['widths'][4]==61,data
  assert data['height']==17 and data['merged'],data
  assert data['a']['fontSize']==14 and data['a']['borderColor']=='#000000',data
  assert data['c']['fontSize']==13 and data['c']['borderColor']=='#000000',data
  assert data['aCss']['font']=='18.67px' and data['cCss']['font']=='17.33px',data
  assert data['aCss']['border']=='rgb(0, 0, 0)' and data['cCss']['border']=='rgb(0, 0, 0)',data
  assert data['persist']==61,data
  save=page.evaluate("""async()=>{
   const bytes=await exportXLSX(book());return Array.from(bytes,v=>String.fromCharCode(v)).join('');
  }""")
  # Reverse the binary string into bytes for independent openpyxl roundtrip.
  reopened=load_workbook(BytesIO(save.encode('latin-1')))
  out=reopened.active
  assert round(out['A1'].font.sz,1)==14
  assert out['A1'].border.top.color.rgb.endswith('000000')
  assert 'A1:D1' in [str(x) for x in out.merged_cells.ranges]
  assert round(out.sheet_format.defaultRowHeight,2)==15
  assert round(out.sheet_format.defaultColWidth,1)==8
  assert not errors,errors
  print('PASS TASK040: Excel default dimensions, narrow widths, border color, font points, merge, qbook, Excel roundtrip')
  browser.close()
