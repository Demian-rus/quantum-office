from pathlib import Path
import io
from shutil import which
from playwright.sync_api import sync_playwright
from openpyxl import load_workbook
root=Path(__file__).resolve().parents[1]
with sync_playwright() as p:
    browser=p.chromium.launch(**({'executable_path':which('chromium')} if which('chromium') else {}),headless=True,args=['--no-sandbox','--disable-dev-shm-usage'])
    page=browser.new_page()
    page.set_content((root/'dist/index.html').read_text(),wait_until='load')
    payload=page.evaluate('''async()=>{
      const b=newWorkbook('roundtrip'),s=currentSheet(b);
      setCell(s,'A1','Название');setCell(s,'B1','Количество');
      setCell(s,'A2','яблоко');setCell(s,'B2','10');
      setCell(s,'A3','груша');setCell(s,'B3','20');
      setCell(s,'A4','яблоко');setCell(s,'B4','15');
      setAutoFilter(s,'A1','B4');setFilterValue(s,0,'яблоко');setFreeze(s,1,1);
      const bytes=await exportXLSX(b);
      const imp=await importXLSX(bytes,'roundtrip.xlsx');
      const r=imp.sheets[0];
      return {bytes:Array.from(bytes),imported:{rules:r.autoFilter?.rules,freeze:r.freeze,visible:[1,2,3].filter(n=>filterRowVisible(r,n)),A3:r.cells.A3.value}};
    }''')
    assert payload['imported']=={'rules':{'0':'яблоко'},'freeze':{'rows':1,'cols':1},'visible':[1,3],'A3':'груша'},payload['imported']
    wb=load_workbook(io.BytesIO(bytes(payload['bytes'])))
    sh=wb.active
    assert sh.auto_filter.ref=='A1:B4',sh.auto_filter.ref
    assert sh.freeze_panes=='B2',sh.freeze_panes
    assert sh['A3'].value=='груша'
    assert sh.row_dimensions[3].hidden is True
    print('PASS XLSX openpyxl + in-app import roundtrip, filter, freeze, hidden rows')
    browser.close()
