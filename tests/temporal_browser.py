"""End-to-end calendar/date UI and Excel numeric-date interchange in Chromium."""
from pathlib import Path
from shutil import which
from io import BytesIO
from datetime import datetime, date, timedelta
from tempfile import TemporaryDirectory
from playwright.sync_api import sync_playwright
from openpyxl import Workbook, load_workbook

root=Path(__file__).resolve().parents[1]
with TemporaryDirectory() as tmp, sync_playwright() as p:
    browser=p.chromium.launch(executable_path=which('chromium') or None,headless=True,args=['--no-sandbox','--disable-dev-shm-usage'])
    page=browser.new_page(viewport={'width':1400,'height':850},accept_downloads=True)
    errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
    page.set_content((root/'dist/index.html').read_text(encoding='utf-8'),wait_until='load')
    page.locator('[data-action=new]').first.click()
    def write(ref,value):
        page.locator('#name-box').fill(ref);page.locator('#name-box').press('Enter')
        field=page.locator('#formula-input');field.fill(value);field.press('Enter')
    def displayed(ref):
        page.locator('#name-box').fill(ref);page.locator('#name-box').press('Enter')
        return page.locator('.grid-cell[data-ref="'+ref+'"]').inner_text()
    write('A1','29.10.2026')
    write('B1','09:15')
    write('B2','11:45')
    write('C2','=B2-B1')
    page.get_by_role('tab',name='Главная').click()
    page.get_by_role('button',name='Длительность').click()
    write('D1','=DATE(2024,2,29)')
    write('D2','=YEAR(D1)')
    write('D3','=MONTH(D1)')
    write('D4','=DAY(D1)')
    assert displayed('A1')=='29.10.2026',displayed('A1')
    assert displayed('B1')=='09:15',displayed('B1')
    assert displayed('C2')=='2:30',displayed('C2')
    assert displayed('D1')=='29.02.2024',displayed('D1')
    assert [displayed('D2'),displayed('D3'),displayed('D4')]==['2024','2','29']
    page.locator('#name-box').fill('A1');page.locator('#name-box').press('Enter')
    page.locator('.grid-cell[data-ref="A1"]').dblclick()
    assert page.locator('.cell-editor').input_value()=='29.10.2026'
    page.locator('.cell-editor').press('Escape')
    # Check XLSX with an independent OOXML reader.
    page.get_by_role('tab',name='Файл').click()
    with page.expect_download() as dl:page.locator('[data-action=xlsx]').first.click()
    sheet=load_workbook(BytesIO(Path(dl.value.path()).read_bytes())).active
    assert isinstance(sheet['A1'].value,(datetime,date)),type(sheet['A1'].value)
    assert sheet['A1'].value.date()==date(2026,10,29)
    assert sheet['B1'].value.hour==9 and sheet['B1'].value.minute==15
    assert sheet['C2'].value=='=B2-B1'
    assert '[h]' in sheet['C2'].number_format.lower(),sheet['C2'].number_format
    assert sheet['D1'].value=='=DATE(2024,2,29)'
    # Import another Excel workbook with date/time styles and assert display.
    imported=Workbook();sh=imported.active;sh.title='Время'
    sh['A1']=datetime(2024,2,29);sh['A1'].number_format='yyyy-mm-dd'
    sh['B1']=datetime(2024,2,29,9,30);sh['B1'].number_format='yyyy-mm-dd hh:mm'
    sh['C1']=timedelta(hours=27,minutes=15);sh['C1'].number_format='[h]:mm'
    target=Path(tmp)/'dates.xlsx';imported.save(target)
    page.locator('#file-picker').set_input_files(str(target))
    # Wait for the *imported workbook*, not an A1 grid cell that already exists
    # in the previous document. Otherwise this check races asynchronous XLSX import.
    page.wait_for_function("() => sheet()?.name === 'Время'")
    assert displayed('A1')=='29.02.2024',displayed('A1')
    assert displayed('B1')=='29.02.2024 09:30',displayed('B1')
    assert displayed('C1')=='27:15',displayed('C1')
    # Explicitly reject an unsupported 1904-epoch source; do not reinterpret it.
    imported.epoch=datetime(1904,1,1)
    older=Path(tmp)/'dates1904.xlsx';imported.save(older)
    page.locator('#file-picker').set_input_files(str(older))
    page.wait_for_timeout(250)
    assert '1904' in page.locator('#toast').inner_text()
    assert displayed('A1')=='29.02.2024'
    assert not errors,errors
    print('PASS temporal UI, formulas, editor, duration, independent XLSX export/import, 1904 guard')
    browser.close()
