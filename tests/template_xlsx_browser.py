from pathlib import Path
from shutil import which
from playwright.sync_api import sync_playwright
from openpyxl import load_workbook
root=Path(__file__).resolve().parents[1]
with sync_playwright() as p:
 browser=p.chromium.launch(headless=True,executable_path=which('chromium') or None,args=['--no-sandbox','--disable-dev-shm-usage'])
 page=browser.new_page(viewport={'width':1440,'height':810},accept_downloads=True)
 page.set_default_timeout(8000)
 page.set_content((root/'dist/index.html').read_text(encoding='utf-8'))
 page.get_by_role('button',name='Шаблоны таблиц').click()
 page.get_by_role('button',name='Недельное расписание занятий').click()
 page.screenshot(path=str(root/'tests/task19-schedule.png'))
 with page.expect_download() as download:
  page.get_by_role('tab',name='Данные').click()
  page.get_by_role('button',name='XLSX').click()
 output=root/'tests/task19-template.xlsx'
 download.value.save_as(str(output))
 wb=load_workbook(output)
 assert wb.sheetnames==['График','Бухгалтерия','Партнёрка'],wb.sheetnames
 assert wb['График']['A4'].value=='Пн'
 assert wb['График']['I2'].value=='Стоимость'
 assert wb['График']['B2'].value=='Ученики'
 assert wb['График']['A4'].fill.fgColor.rgb[-6:]=='FFCA19'
 assert wb['График'].freeze_panes=='C4',wb['График'].freeze_panes
 assert len(wb['График'].merged_cells.ranges)==0
 output.unlink()
 browser.close()
 print('PASS schedule template XLSX export/import readability and explicit no-merge limitation')
