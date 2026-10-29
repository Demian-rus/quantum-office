"""TASK042 — workbook theme, tint and indexed colors survive Excel import."""
from io import BytesIO
from pathlib import Path
from shutil import which
from tempfile import TemporaryDirectory
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET
import base64, colorsys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font, Border, Side, Color
from playwright.sync_api import sync_playwright

root=Path(__file__).resolve().parents[1]
with TemporaryDirectory() as tmp:
 source=Path(tmp)/'theme-source.xlsx'
 updated=Path(tmp)/'theme-customized.xlsx'
 wb=Workbook();ws=wb.active;ws.title='Оформление'
 ws['A1']='Цвет из темы'
 ws['A1'].fill=PatternFill(fill_type='solid',fgColor=Color(theme=4))
 ws['A2']='Светлый'
 ws['A2'].fill=PatternFill(fill_type='solid',fgColor=Color(theme=4,tint=0.5))
 ws['B1']='Индекс'
 ws['B1'].fill=PatternFill(fill_type='solid',fgColor=Color(indexed=10))
 ws['B2']='Шрифт'
 ws['B2'].font=Font(color=Color(theme=1))
 ws['B2'].border=Border(bottom=Side(style='thin',color=Color(theme=4)))
 ws['C1']='Тема LT2'
 ws['C1'].fill=PatternFill(fill_type='solid',fgColor=Color(theme=2))
 ws['C2']='Синий'
 ws['C2'].fill=PatternFill(fill_type='solid',fgColor=Color(rgb='FF00AAFF'))
 wb.save(source)
 ns='{http://schemas.openxmlformats.org/drawingml/2006/main}'
 with ZipFile(source,'r') as zin,ZipFile(updated,'w',ZIP_DEFLATED) as zout:
  for item in zin.infolist():
   data=zin.read(item.filename)
   if item.filename=='xl/theme/theme1.xml':
    theme=ET.fromstring(data);scheme=theme.find('.//'+ns+'clrScheme')
    accent=scheme.find(ns+'accent1')
    assert accent is not None
    inner=list(accent)[0]
    inner.set('lastClr','1260A0') if inner.tag.endswith('sysClr') else inner.set('val','1260A0')
    data=ET.tostring(theme,encoding='utf-8',xml_declaration=True)
   zout.writestr(item,data)
 with sync_playwright() as p:
  browser=p.chromium.launch(executable_path=which('chromium') or None,headless=True,args=['--no-sandbox','--disable-dev-shm-usage'])
  page=browser.new_page(viewport={'width':1400,'height':830})
  page.set_default_timeout(30000)
  errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
  page.set_content((root/'dist/index.html').read_text(encoding='utf-8'),wait_until='load')
  page.locator('#file-picker').set_input_files(str(updated))
  page.wait_for_selector('.grid-cell[data-ref="A1"]')
  state=page.evaluate("""()=>{
   const s=sheet(), refs=['A1','A2','B1','B2','C1','C2'];
   const values=Object.fromEntries(refs.map(ref=>[ref,getEffectiveStyle(s,ref)]));
   const css=Object.fromEntries(refs.map(ref=>{
    const node=document.querySelector('.grid-cell[data-ref="'+ref+'"]');
    const style=getComputedStyle(node);
    return [ref,{bg:style.backgroundColor,color:style.color,border:style.borderBottomColor}];
   }));
   const cloned=decodeBook(encodeBook(book()));
   return {values,css,roundtrip:getEffectiveStyle(cloned.sheets[0],'A1'),
     count:Object.keys(s.cells).length};
  }""")
  assert state['values']['A1']['bg']=='#1260a0',state
  assert state['values']['B1']['bg']=='#ff0000',state
  assert state['values']['B2']['color']=='#000000',state
  assert state['values']['B2']['borderColor']=='#1260a0',state
  assert state['values']['C2']['bg']=='#00aaff',state
  assert state['css']['A1']['bg']=='rgb(18, 96, 160)',state
  assert state['css']['B1']['bg']=='rgb(255, 0, 0)',state
  assert state['roundtrip']['bg']=='#1260a0',state
  # Excel's tint is applied to HLS luminance rather than per RGB channel.
  h,l,sat=colorsys.rgb_to_hls(18/255,96/255,160/255)
  tr,tg,tb=colorsys.hls_to_rgb(h,l*0.5+0.5,sat)
  exp=[round(channel*255) for channel in (tr,tg,tb)]
  got=state['values']['A2']['bg']
  actual=[int(got[i:i+2],16) for i in (1,3,5)]
  assert all(abs(a-b)<=2 for a,b in zip(actual,exp)),(got,exp)
  # Reverse export is RGB; Excel must display the resolved theme color.
  b64=page.evaluate("""async()=>{const bytes=await exportXLSX(book());
   return btoa(Array.from(bytes,v=>String.fromCharCode(v)).join(''));
  }""")
  restored=load_workbook(BytesIO(base64.b64decode(b64)))
  assert restored.active['A1'].fill.fgColor.rgb=='FF1260A0'
  assert restored.active['B1'].fill.fgColor.rgb=='FFFF0000'
  assert restored.active['B2'].border.bottom.color.rgb=='FF1260A0'
  assert state['count']<=6 and not errors,(state,errors)
  print('PASS TASK042: custom theme XML, tint, indexed fill, font, border and OOXML roundtrip')
  browser.close()
