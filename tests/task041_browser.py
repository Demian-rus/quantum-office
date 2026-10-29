"""TASK041 — XLSX row / column borders and merged edges, no phantom fill."""
from pathlib import Path
from tempfile import TemporaryDirectory
from shutil import which
from io import BytesIO
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Border,Side,PatternFill,Font
from playwright.sync_api import sync_playwright

root=Path(__file__).resolve().parents[1]
with TemporaryDirectory() as tmp:
 path=Path(tmp)/'sparse-style-axes.xlsx'
 book=Workbook();sheet=book.active
 black=Side(style='thin',color='FF000000')
 sheet.column_dimensions['E'].width=9
 sheet.column_dimensions['E'].border=Border(bottom=black)
 sheet.row_dimensions[4].border=Border(bottom=black)
 sheet['E3']='Белый'
 sheet['E3'].fill=PatternFill(fill_type='solid',fgColor='FFFFFFFF')
 sheet['E3'].font=Font(bold=True)  # explicit xf without a border
 sheet['H4']='Объединение'
 sheet.merge_cells('H4:I4')
 sheet['A1']='Данные'
 book.save(path)
 with sync_playwright() as p:
  browser=p.chromium.launch(executable_path=which('chromium') or None,headless=True,args=['--no-sandbox','--disable-dev-shm-usage'])
  page=browser.new_page(viewport={'width':1600,'height':900})
  page.set_default_timeout(25000)
  errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
  page.set_content((root/'dist/index.html').read_text(encoding='utf-8'),wait_until='load')
  page.locator('#file-picker').set_input_files(str(path))
  page.wait_for_selector('.grid-cell[data-ref="H4"]')
  state=page.evaluate("""()=>{
   const s=sheet();
   const css=ref=>{const e=document.querySelector('.grid-cell[data-ref="'+ref+'"]');
    const st=getComputedStyle(e);return {bottom:st.borderBottomColor,bottomWidth:st.borderBottomWidth,background:st.backgroundColor};
   };
   return {column:s.columnStyles['4'],row:s.rowStyles?.['3'],
    e2:getEffectiveStyle(s,'E2'),e3:getEffectiveStyle(s,'E3'),
    f4:getEffectiveStyle(s,'F4'),
    e2css:css('E2'),e3css:css('E3'),f4css:css('F4'),h4css:css('H4'),
    cells:Object.keys(s.cells).length,merged:s.merges};
  }""")
  assert state['column']['borders']['bottom'] and state['row']['borders']['bottom'],state
  assert state['e2']['borderColor']=='#000000',state
  assert not state['e3'].get('borders'),state
  assert state['e3']['bg']=='#ffffff',state
  assert state['f4']['borders']['bottom'],state
  assert state['e2css']['bottom']=='rgb(0, 0, 0)',state
  assert state['f4css']['bottom']=='rgb(0, 0, 0)',state
  assert state['h4css']['bottom']=='rgb(0, 0, 0)',state
  assert state['e3css']['background']=='rgb(255, 255, 255)',state
  assert any(x['from']=='H4' and x['to']=='I4' for x in state['merged']),state
  assert state['cells']<10,state
  # Both sparse axis styles survive native QBOOK and real OOXML round-trip.
  encoded=page.evaluate("""async()=>{
   const native=decodeBook(encodeBook(book()));
   if(!getEffectiveStyle(native.sheets[0],'E8').borders?.bottom)throw Error('QBOOK lost column border');
   const out=await exportXLSX(native);return btoa(Array.from(out,c=>String.fromCharCode(c)).join(''));
  }""")
  import base64
  result=load_workbook(BytesIO(base64.b64decode(encoded)))
  assert result.active.column_dimensions['E'].border.bottom.style=='thin'
  assert result.active.row_dimensions[4].border.bottom.style=='thin'
  assert not errors,errors
  print('PASS TASK041: sparse column/row edges, white fill does not erase unrelated borders, merged edge, QBOOK and Excel export')
  browser.close()
