from pathlib import Path
from shutil import which
from tempfile import TemporaryDirectory
from playwright.sync_api import sync_playwright
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

root=Path(__file__).resolve().parents[1]
with TemporaryDirectory() as tmp:
    target=Path(tmp)/'shared_strings_fixture.xlsx'
    book=Workbook(); sh=book.active; sh.title='Бюджет'
    sh['A1']='Имя';sh['B1']='Значение';sh['A2']='Саша';sh['B2']=120;sh['B3']='=B2*2'
    sh['A1'].font=Font(bold=True,color='FF112233');sh['A1'].fill=PatternFill('solid',fgColor='FFCCDDFF');sh['B2'].number_format='#,##0.00'
    sh2=book.create_sheet('Метки'); sh2['A1']='Тест фразы с кириллицей';sh2['A2']='  пробелы  '
    book.save(target)
    # Check workbook uses compressed member(s)
    from zipfile import ZipFile,ZIP_DEFLATED
    with ZipFile(target) as z: assert z.getinfo('xl/worksheets/sheet1.xml').compress_type==ZIP_DEFLATED

    with sync_playwright() as p:
        browser=p.chromium.launch(**({'executable_path':which('chromium')} if which('chromium') else {}),headless=True,args=['--no-sandbox','--disable-dev-shm-usage'])
        context=browser.new_context(accept_downloads=True,viewport={'width':1400,'height':900})
        page=context.new_page();errors=[]
        page.on('pageerror',lambda e:errors.append(str(e)))
        page.set_content((root/'dist/index.html').read_text(encoding='utf-8'),wait_until='load')
        page.locator('#file-picker').set_input_files(str(target))
        page.wait_for_selector('.grid-cell[data-ref="A1"]')
        assert page.locator('#book-name').input_value()=='shared_strings_fixture'
        assert page.locator('.grid-cell[data-ref="A2"]').inner_text()=='Саша'
        assert page.locator('.grid-cell[data-ref="B2"]').inner_text()=='120'
        assert page.locator('.grid-cell[data-ref="B3"]').inner_text()=='240'
        page.locator('#name-box').fill('B3');page.locator('#name-box').press('Enter');assert page.locator('#formula-input').input_value()=='=B2*2'
        page.locator('.sheet-tab').last.click()
        assert page.locator('.grid-cell[data-ref="A1"]').inner_text()=='Тест фразы с кириллицей'
        page.get_by_role('tab',name='Данные').click()
        with page.expect_download() as event:
            page.get_by_role('button',name='XLSX').click()
        download=event.value
        export_target=Path(tmp)/download.suggested_filename;download.save_as(export_target)
        exported=load_workbook(export_target)
        assert exported.sheetnames==['Бюджет','Метки'],exported.sheetnames
        assert exported['Бюджет']['B3'].value=='=B2*2'
        assert exported['Бюджет']['A2'].value=='Саша'
        assert exported['Бюджет']['A1'].font.bold
        assert exported['Бюджет']['A1'].fill.fgColor.rgb[-6:]=='CCDDFF'
        assert exported['Метки']['A2'].value=='  пробелы  '
        # Reject invalid XLSX without altering workspace
        bad=Path(tmp)/'bad.xlsx';bad.write_bytes(b'not a zip')
        page.locator('#file-picker').set_input_files(str(bad))
        page.wait_for_timeout(500)
        assert page.locator('.sheet-tab').count()==2
        assert not errors,errors
        print('XLSX browser PASS: deflated import, formulas, formatting, multi-sheet export, openpyxl re-open, malformed archive rejection')
        browser.close()
