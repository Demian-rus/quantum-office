"""TASK043 — Excel mixed border weights, colors and merged outline."""
from pathlib import Path
from tempfile import TemporaryDirectory
from io import BytesIO
from shutil import which
import base64
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Side, PatternFill
from playwright.sync_api import sync_playwright

root=Path(__file__).resolve().parents[1]
with TemporaryDirectory() as tmp:
 path=Path(tmp)/'border-mix.xlsx'
 wb=Workbook();ws=wb.active
 ws.title='Оформление'
 ws['B2']='Четыре стороны'
 ws['B2'].border=Border(
  left=Side(style='medium',color='FF00AA00'),
  right=Side(style='double',color='FF0000AA'),
  top=Side(style='thick',color='FFAA0000'),
  bottom=Side(style='dotted',color='FFAA00AA'))
 ws['B2'].fill=PatternFill(fill_type='solid',fgColor='FFFFFFFF')
 ws['D3']='Соседняя'
 ws['D3'].border=Border(bottom=Side(style='medium',color='FF0033AA'))
 ws['F4']='Объединение'
 ws.merge_cells('F4:G4')
 ws['F4'].border=Border(left=Side(style='medium',color='FF005500'),
                        top=Side(style='dashed',color='FFAA5500'))
 ws['G4'].border=Border(right=Side(style='thick',color='FFAA00AA'))
 wb.save(path)
 with sync_playwright() as p:
  browser=p.chromium.launch(executable_path=which('chromium') or None,headless=True,args=['--no-sandbox','--disable-dev-shm-usage'])
  page=browser.new_page(viewport={'width':1500,'height':840})
  page.set_default_timeout(25000)
  errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
  page.set_content((root/'dist/index.html').read_text(encoding='utf-8'),wait_until='load')
  page.locator('#file-picker').set_input_files(str(path))
  page.wait_for_selector('.grid-cell[data-ref="B2"]')
  result=page.evaluate("""()=>{
   const sheet0=sheet();
   const css=ref=>{
    const el=document.querySelector('.grid-cell[data-ref="'+ref+'"]');
    const st=getComputedStyle(el);
    return Object.fromEntries(['top','right','bottom','left'].map(side=>{
     const k=side[0].toUpperCase()+side.slice(1);
     return [side,{width:st['border'+k+'Width'],style:st['border'+k+'Style'],color:st['border'+k+'Color']}];
    }));
   };
   const native=decodeBook(encodeBook(book()));
   return {b2:getEffectiveStyle(sheet0,'B2'),d3:getEffectiveStyle(sheet0,'D3'),
    b2css:css('B2'),d3css:css('D3'),merge:css('F4'),
    merged:sheet0.merges.some(m=>m.from==='F4'&&m.to==='G4'),
    round:getEffectiveStyle(native.sheets[0],'B2'),
    populated:Object.keys(sheet0.cells).length};
  }""")
  assert result['b2']['borderEdges']['top']=={'line':'thick','color':'#aa0000'},result
  assert result['b2']['borderEdges']['right']=={'line':'double','color':'#0000aa'},result
  assert result['b2css']['top']=={'width':'3px','style':'solid','color':'rgb(170, 0, 0)'},result
  assert result['b2css']['right']=={'width':'3px','style':'double','color':'rgb(0, 0, 170)'},result
  assert result['b2css']['bottom']=={'width':'1px','style':'dotted','color':'rgb(170, 0, 170)'},result
  assert result['b2css']['left']=={'width':'2px','style':'solid','color':'rgb(0, 170, 0)'},result
  assert result['d3css']['bottom']['width']=='2px',result
  assert result['d3css']['bottom']['color']=='rgb(0, 51, 170)',result
  assert result['merge']['left']['color']=='rgb(0, 85, 0)',result
  assert result['merge']['right']['color']=='rgb(170, 0, 170)',result
  assert result['merged'] and result['round']['borderEdges']==result['b2']['borderEdges'],result
  assert result['populated']<16 and not errors,(result,errors)
  binary=page.evaluate("""async()=>{const bytes=await exportXLSX(book());
   return btoa(Array.from(bytes,v=>String.fromCharCode(v)).join(''));
  }""")
  reopened=load_workbook(BytesIO(base64.b64decode(binary)))
  cell=reopened.active['B2']
  assert cell.border.left.style=='medium' and cell.border.left.color.rgb=='FF00AA00'
  assert cell.border.right.style=='double' and cell.border.right.color.rgb=='FF0000AA'
  assert cell.border.top.style=='thick' and cell.border.top.color.rgb=='FFAA0000'
  assert cell.border.bottom.style=='dotted' and cell.border.bottom.color.rgb=='FFAA00AA'
  assert reopened.active['D3'].border.bottom.style=='medium'
  assert reopened.active['F4'].border.left.style=='medium'
  print('PASS TASK043: mixed Excel side styles, CSS display, QBOOK and XLSX roundtrip')
  browser.close()
